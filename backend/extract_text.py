"""Turn an uploaded lease into plain text without sending it anywhere.

PDFs, Word files and .txt files are read directly. Photos, and scanned PDFs with no
text layer, go through RapidOCR, which runs locally on this machine.
"""
import io
import os

import numpy as np
from docx import Document
from PIL import Image, ImageOps
from pypdf import PdfReader

IMAGE_TYPES = (".png", ".jpg", ".jpeg", ".webp")
SUPPORTED = (".pdf", ".docx", ".txt") + IMAGE_TYPES
MAX_SIDE = 2000  # shrink big phone photos so OCR stays fast

_ocr = None


def _get_ocr():
    global _ocr
    if _ocr is None:
        from rapidocr_onnxruntime import RapidOCR

        _ocr = RapidOCR()
    return _ocr


def _ocr_image(data):
    img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    img.thumbnail((MAX_SIDE, MAX_SIDE))
    result, _ = _get_ocr()(np.array(img))
    return "\n".join(line[1] for line in result or [])


def extract_text(filename, data):
    """Return (text, used_ocr)."""
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages).strip()
        if text:
            return text, False
        # No text layer: probably a scan, so OCR the page images.
        pages = []
        for page in reader.pages:
            pages.extend(_ocr_image(image.data) for image in page.images)
        return "\n\n".join(pages).strip(), True
    if ext == ".docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip(), False
    if ext == ".txt":
        return data.decode("utf-8", errors="replace").strip(), False
    if ext in IMAGE_TYPES:
        return _ocr_image(data).strip(), True
    raise ValueError(f"Unsupported file type. Please upload one of: {', '.join(SUPPORTED)}")
